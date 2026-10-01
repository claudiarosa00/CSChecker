"""
Lê data/base-formatos.xlsx (a mesma base usada pelo CSBuilder) e gera
data/specs.json: uma lista de formatos com as regras de validação já
estruturadas (dimensões, peso máximo, tipos de ficheiro, aspect ratio),
em vez do texto livre original — para o app.js poder verificar
materiais sem ter de interpretar texto a correr no browser.
"""
import json
import os
import re

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEM = os.path.join(RAIZ, "data", "base-formatos.xlsx")
DESTINO = os.path.join(RAIZ, "data", "specs.json")

MAPA_MEIO = {
    "INTERNET": "Digital",
    "PROGRAMÁTICO": "Digital",
    "OOH": "OOH",
    "TV": "TV",
    "RÁDIO": "Rádio",
    "RADIO": "Rádio",
    "CINEMA": "Cinema",
    "IMPRENSA": "Imprensa",
}

RE_DIMENSAO = re.compile(r"(\d+)\s*[xX]\s*(\d+)")
RE_NUMERO_UNIDADE = re.compile(r"([\d.,]+)\s*(KB|MB|GB)", re.IGNORECASE)
UNIDADE_PARA_BYTES = {"KB": 1024, "MB": 1024 * 1024, "GB": 1024 * 1024 * 1024}


def dividir_opcoes(texto):
    """'970x250, 320x10' ou '1240x150 ou 1240x120, 320x100' -> lista de blocos."""
    if not texto:
        return []
    texto = texto.replace(" ou ", ", ").replace(" or ", ", ")
    return [parte.strip() for parte in texto.split(",") if parte.strip()]


def extrair_dimensoes(texto):
    opcoes = []
    for bloco in dividir_opcoes(texto):
        m = RE_DIMENSAO.search(bloco)
        if m:
            opcoes.append({"largura": int(m.group(1)), "altura": int(m.group(2))})
    return opcoes


def extrair_pesos_bytes(texto):
    """Pode haver vários pesos (um por opção de dimensão) — devolve a lista
    na mesma ordem em que aparecem, em bytes."""
    if not texto:
        return []
    pesos = []
    for bloco in dividir_opcoes(texto):
        m = RE_NUMERO_UNIDADE.search(bloco)
        if m:
            valor = float(m.group(1).replace(",", "."))
            unidade = m.group(2).upper()
            pesos.append(round(valor * UNIDADE_PARA_BYTES[unidade]))
    return pesos


def extrair_tipos_ficheiro(texto):
    if not texto:
        return []
    tipos = dividir_opcoes(texto)
    # normaliza (maiúsculas, sem espaços) e separa "JPG/PNG" se vier junto
    resultado = []
    for tipo in tipos:
        for sub in re.split(r"[/]", tipo):
            sub = sub.strip().upper()
            if sub and sub not in resultado:
                resultado.append(sub)
    return resultado


def extrair_aspect_ratios(texto):
    if not texto:
        return []
    resultado = []
    for bloco in dividir_opcoes(texto):
        m = re.search(r"(\d+)\s*:\s*(\d+)", bloco)
        if m:
            largura, altura = int(m.group(1)), int(m.group(2))
            resultado.append({"largura": largura, "altura": altura, "valor": largura / altura})
    return resultado


def associar_peso_a_dimensao(dimensoes, pesos):
    """Quando o número de pesos bate certo com o de dimensões, cada
    dimensão leva o seu próprio limite; caso contrário, aplica-se a
    todas o maior peso listado (mais permissivo, para não recusar um
    material válido só por não sabermos emparelhar)."""
    if not dimensoes:
        return None
    if len(pesos) == len(dimensoes):
        for dim, peso in zip(dimensoes, pesos):
            dim["pesoMaximoBytes"] = peso
        return None
    return max(pesos) if pesos else None


def main():
    wb = openpyxl.load_workbook(ORIGEM, data_only=True)
    ws = wb.active
    cabecalhos = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(cabecalhos)}

    formatos = []
    for linha_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        meio_bruto = (row[idx["Meio"]] or "").strip().upper()
        meio = MAPA_MEIO.get(meio_bruto)
        if not meio:
            continue

        fornecedor = row[idx["Fornecedor"]] or ""
        veiculo = row[idx["Veículo"]] or fornecedor
        nome_formato = row[idx["Formato"]] or ""
        if not nome_formato:
            continue

        dimensao_texto = row[idx["Dimensão"]] or ""
        peso_texto = row[idx["Peso"]] or ""
        tipo_texto = row[idx["Tipo de Ficheiro"]] or ""
        aspect_texto = row[idx["Aspect Ratio"]] or ""

        dimensoes = extrair_dimensoes(dimensao_texto)
        pesos = extrair_pesos_bytes(peso_texto)
        peso_maximo_geral = associar_peso_a_dimensao(dimensoes, pesos)

        formatos.append({
            "id": f"f{linha_num}",
            "meio": meio,
            "fornecedor": fornecedor,
            "veiculo": veiculo,
            "formato": nome_formato,
            "dimensaoTexto": dimensao_texto,
            "pesoTexto": peso_texto,
            "tipoFicheiroTexto": tipo_texto,
            "aspectRatioTexto": aspect_texto,
            "dimensoes": dimensoes,
            "pesoMaximoBytes": peso_maximo_geral,
            "tiposFicheiro": extrair_tipos_ficheiro(tipo_texto),
            "aspectRatios": extrair_aspect_ratios(aspect_texto),
        })

    os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
    with open(DESTINO, "w", encoding="utf-8") as f:
        json.dump(formatos, f, ensure_ascii=False, indent=1)

    print(f"{len(formatos)} formatos escritos em {DESTINO}")
    com_dimensoes = sum(1 for f in formatos if f["dimensoes"])
    print(f"{com_dimensoes} com pelo menos uma dimensão estruturada")


if __name__ == "__main__":
    main()
